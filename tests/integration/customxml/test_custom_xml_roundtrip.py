"""Round-trip tests for packages containing multiple Custom XML parts."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Protocol, cast
from zipfile import ZipFile

import pytest

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.parser import parse_xml
from docx.parts.customxml import CustomXmlPart

FIXTURE = (
    Path(__file__).parent.parent.parent
    / "fixtures"
    / "customxml"
    / "multiple-custom-xml-parts.docx"
)
CUSTOMERS_NS = "urn:python-docx-oss:customers"
INVENTORY_NS = "urn:python-docx-oss:inventory"
AUDIT_NS = "urn:python-docx-oss:audit"


class _CustomXmlPartProtocol(Protocol):
    def add_item(self, tag: str, text: str = "", **attrib: str) -> None: ...


def _parts_by_name(document: DocumentType) -> dict[str, CustomXmlPart]:
    return {str(part.partname): part for part in document.part.custom_xml_parts}


class DescribeMultipleCustomXmlParts:
    def it_loads_multiple_namespaced_parts_with_non_contiguous_names(self):
        document = Document(io.BytesIO(FIXTURE.read_bytes()))
        parts = _parts_by_name(document)

        assert set(parts) == {"/customXml/item1.xml", "/customXml/item3.xml"}
        assert parts["/customXml/item1.xml"].tag == f"{{{CUSTOMERS_NS}}}customers"
        assert parts["/customXml/item3.xml"].tag == f"{{{INVENTORY_NS}}}inventory"
        assert len(parts["/customXml/item1.xml"].items) == 2
        assert len(parts["/customXml/item3.xml"].items) == 2

    def it_round_trips_create_update_and_delete_operations_in_memory(self):
        document = Document(io.BytesIO(FIXTURE.read_bytes()))
        parts = _parts_by_name(document)
        customers = parts["/customXml/item1.xml"]
        inventory = parts["/customXml/item3.xml"]

        cast(_CustomXmlPartProtocol, customers).add_item(
            f"{{{CUSTOMERS_NS}}}status",
            "active",
            source="round-trip",
        )
        inventory.delete_item(inventory.items[0])
        document.part.add_custom_xml_part(
            f'<audit xmlns="{AUDIT_NS}" revision="1"><event>created</event></audit>',
        )

        output = io.BytesIO()
        document.save(output)
        reopened = Document(io.BytesIO(output.getvalue()))
        reopened_parts = _parts_by_name(reopened)

        assert set(reopened_parts) == {
            "/customXml/item1.xml",
            "/customXml/item2.xml",
            "/customXml/item3.xml",
        }
        status = reopened_parts["/customXml/item1.xml"].items[-1]
        assert status.tag == f"{{{CUSTOMERS_NS}}}status"
        assert status.text == "active"
        assert status.attrib["source"] == "round-trip"
        assert len(reopened_parts["/customXml/item3.xml"].items) == 1
        assert reopened_parts["/customXml/item2.xml"].tag == f"{{{AUDIT_NS}}}audit"

        with ZipFile(io.BytesIO(output.getvalue())) as package:
            names = set(package.namelist())
        assert "customXml/itemProps1.xml" in names
        assert "customXml/itemProps3.xml" in names
        assert "customXml/_rels/item1.xml.rels" in names
        assert "customXml/_rels/item3.xml.rels" in names

    def it_round_trips_replacing_a_custom_xml_part_element(self):
        document = Document(io.BytesIO(FIXTURE.read_bytes()))
        customers = _parts_by_name(document)["/customXml/item1.xml"]
        wrapper = parse_xml('<wrapper xmlns="urn:python-docx-oss:wrapper" version="1"/>')
        wrapper.append(customers.element)
        customers.element = wrapper

        output = io.BytesIO()
        document.save(output)
        reopened = Document(io.BytesIO(output.getvalue()))
        replaced = _parts_by_name(reopened)["/customXml/item1.xml"]

        assert replaced.tag == "{urn:python-docx-oss:wrapper}wrapper"
        assert replaced.attrib["version"] == "1"  # pyright: ignore[reportUnknownMemberType]
        assert len(replaced.items) == 1
        assert replaced.items[0].tag == f"{{{CUSTOMERS_NS}}}customers"
        assert len(replaced.items[0]) == 2
        with ZipFile(io.BytesIO(output.getvalue())) as package:
            names = set(package.namelist())
        assert "customXml/itemProps1.xml" in names
        assert "customXml/_rels/item1.xml.rels" in names

    def it_rejects_an_explicit_partname_collision(self):
        document = Document(io.BytesIO(FIXTURE.read_bytes()))

        with pytest.raises(ValueError, match="part already exists at '/customXml/item1.xml'"):
            document.part.add_custom_xml_part("<duplicate/>", file_name="item1")

    def it_round_trips_deleting_a_part_and_reuses_its_partname(self):
        document = Document(io.BytesIO(FIXTURE.read_bytes()))
        customers = _parts_by_name(document)["/customXml/item1.xml"]

        document.part.delete_custom_xml_part(customers)
        replacement = document.part.add_custom_xml_part(f'<audit xmlns="{AUDIT_NS}"/>')
        assert str(replacement.partname) == "/customXml/item1.xml"

        output = io.BytesIO()
        document.save(output)
        reopened = Document(io.BytesIO(output.getvalue()))
        reopened_parts = _parts_by_name(reopened)

        assert set(reopened_parts) == {"/customXml/item1.xml", "/customXml/item3.xml"}
        assert reopened_parts["/customXml/item1.xml"].tag == f"{{{AUDIT_NS}}}audit"
        with ZipFile(io.BytesIO(output.getvalue())) as package:
            names = set(package.namelist())
        assert "customXml/itemProps1.xml" not in names
        assert "customXml/_rels/item1.xml.rels" not in names
        assert "customXml/itemProps3.xml" in names
        assert "customXml/_rels/item3.xml.rels" in names

    def it_rejects_deleting_a_part_twice(self):
        document = Document(io.BytesIO(FIXTURE.read_bytes()))
        customers = _parts_by_name(document)["/customXml/item1.xml"]
        document.part.delete_custom_xml_part(customers)

        with pytest.raises(ValueError, match="is not related to this document"):
            document.part.delete_custom_xml_part(customers)
