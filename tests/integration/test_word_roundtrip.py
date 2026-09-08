"""Round-trip tests for documents produced by Microsoft Word."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from lxml import etree

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FIXTURE_DIR = Path(__file__).parent.parent / "fixtures"
TEST_FILE_DIR = Path(__file__).parent.parent / "test_files"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _fixture(*parts: str) -> Path:
    return FIXTURE_DIR.joinpath(*parts)


def _round_trip(path: Path) -> dict[str, bytes]:
    document = Document(io.BytesIO(path.read_bytes()))
    document.add_paragraph("Round-trip verification")

    output = io.BytesIO()
    document.save(output)

    reopened = Document(io.BytesIO(output.getvalue()))
    assert reopened.paragraphs[-1].text == "Round-trip verification"

    with ZipFile(io.BytesIO(output.getvalue())) as package:
        return {name: package.read(name) for name in package.namelist()}


def _xml(package: dict[str, bytes], partname: str) -> Any:
    return etree.fromstring(package[partname])


def _relationship_target(package: dict[str, bytes], relationship_id: str) -> str:
    relationships = _xml(package, "word/_rels/document.xml.rels")
    relationship = relationships.find(f"{{{PR_NS}}}Relationship[@Id='{relationship_id}']")
    assert relationship is not None
    target = relationship.get("Target")
    assert target is not None
    return target


def _content_type(package: dict[str, bytes], partname: str) -> str | None:
    content_types = _xml(package, "[Content_Types].xml")
    override = content_types.find(f"{{{CT_NS}}}Override[@PartName='{partname}']")
    if override is not None:
        return override.get("ContentType")

    extension = partname.rsplit(".", 1)[-1]
    default = content_types.find(f"{{{CT_NS}}}Default[@Extension='{extension}']")
    return None if default is None else default.get("ContentType")


class DescribeWordDocumentRoundTrips:
    def it_round_trips_a_minimal_word_document_in_memory(self):
        package = _round_trip(_fixture("documents", "word-blank.docx"))

        assert "word/document.xml" in package
        assert "word/styles.xml" in package
        assert not any(name.startswith("word/media/") for name in package)

    def it_round_trips_a_default_numbering_part_when_the_source_has_none(self):
        fixture = _fixture("documents", "word-blank.docx")
        with ZipFile(fixture) as source_package:
            assert "word/numbering.xml" not in source_package.namelist()

        document = Document(io.BytesIO(fixture.read_bytes()))
        numbering_part = document.part.numbering_part
        output = document.to_bytes()
        reopened = Document(output)

        assert str(numbering_part.partname) == "/word/numbering.xml"
        assert str(reopened.part.numbering_part.partname) == "/word/numbering.xml"
        with ZipFile(io.BytesIO(output)) as package_file:
            package = {name: package_file.read(name) for name in package_file.namelist()}
        assert "word/numbering.xml" in package
        assert _content_type(package, "/word/numbering.xml") == CT.WML_NUMBERING
        relationships = _xml(package, "word/_rels/document.xml.rels")
        numbering_relationship = relationships.find(
            f"{{{PR_NS}}}Relationship[@Type='{RT.NUMBERING}']"
        )
        assert numbering_relationship is not None
        assert numbering_relationship.get("Target") == "numbering.xml"

    def it_round_trips_a_changed_external_hyperlink_target(self):
        document = Document(_fixture("documents", "word-blank.docx").read_bytes())
        paragraph = document.add_paragraph()
        relationship_id = document.part.relate_to(
            "https://example.invalid/objects/old-id",
            RT.HYPERLINK,
            is_external=True,
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship_id)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Business link"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)  # pyright: ignore[reportPrivateUsage]

        new_target = "https://example.invalid/objects/new-id?position=begin"
        document.part.rels[relationship_id].target_ref = new_target
        output = document.to_bytes()
        reopened = Document(output)

        assert reopened.paragraphs[-1].hyperlinks[0].text == "Business link"
        assert reopened.paragraphs[-1].hyperlinks[0].address == new_target
        with ZipFile(io.BytesIO(output)) as package_file:
            relationships = etree.fromstring(package_file.read("word/_rels/document.xml.rels"))
        relationship = relationships.find(f"{{{PR_NS}}}Relationship[@Id='{relationship_id}']")
        assert relationship is not None
        assert relationship.get("Target") == new_target
        assert relationship.get("TargetMode") == "External"

    def it_creates_independent_document_copies_from_one_snapshot(self):
        source = Document(_fixture("customxml", "multiple-custom-xml-parts.docx").read_bytes())
        source.add_paragraph("Shared snapshot")

        document_b = source.copy()
        document_c = source.copy()
        document_d = source.copy()
        document_b.add_paragraph("Version B")
        document_c.add_paragraph("Version C")
        document_d.part.custom_xml_parts[0].add_item("version", "D")

        assert source.part.package is not document_b.part.package
        assert document_b.part.package is not document_c.part.package
        assert document_c.part.package is not document_d.part.package
        assert source.paragraphs[-1].text == "Shared snapshot"
        assert document_b.paragraphs[-1].text == "Version B"
        assert document_c.paragraphs[-1].text == "Version C"
        assert document_d.paragraphs[-1].text == "Shared snapshot"
        assert len(source.part.custom_xml_parts[0].items) == 2
        assert len(document_d.part.custom_xml_parts[0].items) == 3

        reopened = Document(document_d.to_bytes())
        assert reopened.paragraphs[-1].text == "Shared snapshot"
        assert reopened.part.custom_xml_parts[0].items[-1].text == "D"

    def it_preserves_a_floating_picture(self):
        package = _round_trip(_fixture("drawing", "word-floating-picture.docx"))
        document = _xml(package, "word/document.xml")

        anchor = document.find(f".//{{{WP_NS}}}anchor")
        assert anchor is not None
        assert anchor.get("behindDoc") == "1"
        position_h = anchor.find(f"{{{WP_NS}}}positionH")
        position_v = anchor.find(f"{{{WP_NS}}}positionV")
        assert position_h is not None
        assert position_v is not None
        assert position_h.get("relativeFrom") == "column"
        assert position_v.get("relativeFrom") == "paragraph"
        assert anchor.find(f"{{{WP_NS}}}wrapTight") is not None

        blip = anchor.find(f".//{{{A_NS}}}blip")
        assert blip is not None
        relationship_id = blip.get(f"{{{R_NS}}}embed")
        assert relationship_id is not None
        target = _relationship_target(package, relationship_id)
        assert f"word/{target}" in package

    def it_preserves_an_svg_image_and_its_png_fallback(self):
        package = _round_trip(_fixture("images", "word-svg-image.docx"))
        document = _xml(package, "word/document.xml")

        blip = document.find(f".//{{{A_NS}}}blip")
        svg_blip = document.find(f".//{{{ASVG_NS}}}svgBlip")
        assert blip is not None
        assert svg_blip is not None

        png_relationship_id = blip.get(f"{{{R_NS}}}embed")
        svg_relationship_id = svg_blip.get(f"{{{R_NS}}}embed")
        assert png_relationship_id is not None
        assert svg_relationship_id is not None
        png_target = _relationship_target(package, png_relationship_id)
        svg_target = _relationship_target(package, svg_relationship_id)
        assert png_target.endswith(".png")
        assert svg_target.endswith(".svg")
        assert f"word/{png_target}" in package
        assert f"word/{svg_target}" in package
        assert _content_type(package, f"/word/{svg_target}") == "image/svg+xml"

    def it_creates_and_round_trips_an_emf_image_in_memory(self):
        emf_blob = (TEST_FILE_DIR / "sonic.emf").read_bytes()
        document = Document(io.BytesIO(_fixture("documents", "word-blank.docx").read_bytes()))
        document.add_picture(io.BytesIO(emf_blob))

        output = io.BytesIO()
        document.save(output)
        reopened = Document(io.BytesIO(output.getvalue()))
        assert len(reopened.inline_shapes) == 1

        round_trip_output = io.BytesIO()
        reopened.save(round_trip_output)
        with ZipFile(io.BytesIO(round_trip_output.getvalue())) as docx_package:
            package = {name: docx_package.read(name) for name in docx_package.namelist()}

        document_xml = _xml(package, "word/document.xml")
        blip = document_xml.find(f".//{{{A_NS}}}blip")
        assert blip is not None
        relationship_id = blip.get(f"{{{R_NS}}}embed")
        assert relationship_id is not None
        emf_target = _relationship_target(package, relationship_id)
        assert emf_target.endswith(".emf")
        assert package[f"word/{emf_target}"] == emf_blob
        assert _content_type(package, f"/word/{emf_target}") == "image/x-emf"

    def it_preserves_distinct_latin_and_east_asian_fonts(self):
        package = _round_trip(_fixture("fonts", "word-east-asian-fonts.docx"))
        document = _xml(package, "word/document.xml")

        chinese_run = next(
            run
            for run in document.findall(f".//{{{W_NS}}}r")
            if (text := run.find(f"{{{W_NS}}}t")) is not None and "全局笔记" in (text.text or "")
        )
        fonts = chinese_run.find(f"{{{W_NS}}}rPr/{{{W_NS}}}rFonts")
        assert fonts is not None
        assert fonts.get(f"{{{W_NS}}}ascii") == "Open Sans"
        assert fonts.get(f"{{{W_NS}}}hAnsi") == "Open Sans"
        assert fonts.get(f"{{{W_NS}}}eastAsia") == "宋体"
