"""Hyperlink normalization preserves text, formatting, selection, and package data."""

from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree
from lxml.etree import _Element as etree_Element  # pyright: ignore[reportPrivateUsage]

from docx import Document, inspect, normalize
from docx.transforms import HyperlinkTarget
from tests.inspection.helpers import CT, RT, R, W, package

from .helpers import REL_ATTR, complex_field, document_parts, link_runs, unzip


def _root(content: bytes, name: str = "word/document.xml") -> etree_Element:
    return etree.fromstring(unzip(content)[name])


class DescribeHyperlinkFieldConversion:
    def it_converts_complex_fields_preserving_result_run_formatting_and_text(self):
        content = package(document_parts(complex_field()))
        result = normalize(content, hyperlink_fields=True)
        assert result.report.succeeded
        assert result.report.changed
        root = _root(result.content)
        links = list(root.iter(f"{{{W}}}hyperlink"))
        assert len(links) == 1
        assert not list(root.iter(f"{{{W}}}fldChar"))
        assert not list(root.iter(f"{{{W}}}instrText"))
        runs = links[0].findall(f"{{{W}}}r")
        assert len(runs) == 2
        assert runs[0].find(f"{{{W}}}rPr/{{{W}}}b") is not None
        assert runs[1].find(f"{{{W}}}rPr/{{{W}}}i") is not None
        assert "".join(node.text or "" for node in links[0].iter(f"{{{W}}}t")) == "Hello world"
        rel_id = links[0].get(f"{{{REL_ATTR}}}id")
        rels = _root(result.content, "word/_rels/document.xml.rels")
        assert any(
            rel.get("Id") == rel_id and rel.get("Target") == "https://example.test/doc"
            for rel in rels
        )
        assert Document(result.content).paragraphs[0].hyperlinks[0].text == "Hello world"
        again = normalize(result.content, hyperlink_fields=True)
        assert again.report.succeeded
        assert not again.report.changed
        assert again.content == result.content

    def it_creates_missing_relationship_parts_and_preserves_existing_ids(self):
        parts = document_parts(complex_field())
        del parts["word/_rels/document.xml.rels"]
        result = normalize(package(parts), hyperlink_fields=True)
        assert result.report.succeeded
        assert "word/_rels/document.xml.rels" in unzip(result.content)
        assert len(_root(result.content, "word/_rels/document.xml.rels")) == 1
        with ZipFile(io.BytesIO(result.content)) as archive:
            assert archive.getinfo("word/_rels/document.xml.rels").compress_type == ZIP_DEFLATED

    def it_reuses_existing_external_hyperlink_relationships(self):
        parts = document_parts(complex_field())
        rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
        etree.SubElement(
            rels,
            f"{{{R}}}Relationship",
            Id="external42",
            Type=RT + "hyperlink",
            Target="https://example.test/doc",
            TargetMode="External",
        )
        parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
        result = normalize(package(parts), hyperlink_fields=True)
        assert result.report.succeeded
        link = _root(result.content).find(f".//{{{W}}}hyperlink")
        assert link is not None
        assert link.get(f"{{{REL_ATTR}}}id") == "external42"
        assert not any(
            change.code == "added_hyperlink_relationship" for change in result.report.changes
        )

    def it_reads_instructions_split_across_runs(self):
        field = complex_field().replace(
            "HYPERLINK", "HYPER</w:instrText></w:r><w:r><w:instrText>LINK"
        )
        result = normalize(package(document_parts(field)), hyperlink_fields=True)
        assert result.report.succeeded
        assert result.report.changed

    @pytest.mark.parametrize(
        ("instruction", "anchor", "address"),
        [
            ('HYPERLINK \\l "Chapter1"', "Chapter1", None),
            ('HYPERLINK "https://example.test" \\l "Chapter1"', "Chapter1", "https://example.test"),
            ("HYPERLINK https://example.test", None, "https://example.test"),
        ],
    )
    def it_preserves_external_and_bookmark_targets(
        self, instruction: str, anchor: str | None, address: str | None
    ):
        result = normalize(
            package(document_parts(complex_field(instruction))), hyperlink_fields=True
        )
        assert result.report.succeeded
        link = next(_root(result.content).iter(f"{{{W}}}hyperlink"))
        assert link.get(f"{{{W}}}anchor") == anchor
        if address is None:
            assert link.get(f"{{{REL_ATTR}}}id") is None

    def it_preserves_tooltip_and_target_frame(self):
        field = complex_field(
            'HYPERLINK "https://example.test" \\o "A tooltip" \\t "_blank" \\* MERGEFORMAT'
        )
        result = normalize(package(document_parts(field)), hyperlink_fields=True)
        assert result.report.succeeded
        link = next(_root(result.content).iter(f"{{{W}}}hyperlink"))
        assert link.get(f"{{{W}}}tooltip") == "A tooltip"
        assert link.get(f"{{{W}}}tgtFrame") == "_blank"

    def it_converts_simple_fields(self):
        body = "<w:p><w:fldSimple w:instr='HYPERLINK \"https://example.test\"'><w:r><w:t>label</w:t></w:r></w:fldSimple></w:p>"
        result = normalize(package(document_parts(body)), hyperlink_fields=True)
        assert result.report.succeeded
        assert result.report.changed
        assert not list(_root(result.content).iter(f"{{{W}}}fldSimple"))
        assert Document(result.content).paragraphs[0].hyperlinks[0].text == "label"

    def it_allows_simple_fields_through_the_inspection_precheck(self):
        body = (
            "<w:p><w:fldSimple w:instr='HYPERLINK \"https://example.test\"'>"
            "<w:r><w:t>label</w:t></w:r></w:fldSimple></w:p>"
        )
        content = package(document_parts(body))
        assert inspect(content).features.field_hyperlink
        result = normalize(content, hyperlink_fields=True)
        assert result.report.changed
        assert not inspect(result.content).features.field_hyperlink

    @pytest.mark.parametrize(
        ("instruction", "expected"),
        [
            (r'HYPERLINK "C:\\docs\\a.docx"', "C:\\docs\\a.docx"),
            (r'HYPERLINK "\\\\server\\share\\a.docx"', "\\\\server\\share\\a.docx"),
            (r'HYPERLINK "https://example.test/a\"b"', 'https://example.test/a"b'),
        ],
    )
    @pytest.mark.parametrize("simple", [False, True])
    def it_decodes_word_quoted_string_escapes(
        self,
        instruction: str,
        expected: str,
        simple: bool,
    ):
        if simple:
            escaped = instruction.replace('"', "&quot;")
            body = (
                f'<w:p><w:fldSimple w:instr="{escaped}">'
                "<w:r><w:t>label</w:t></w:r></w:fldSimple></w:p>"
            )
        else:
            body = complex_field(instruction)
        result = normalize(package(document_parts(body)), hyperlink_fields=True)
        assert result.report.succeeded
        assert result.report.changed
        rels = _root(result.content, "word/_rels/document.xml.rels")
        assert any(
            rel.get("Type") == RT + "hyperlink" and rel.get("Target") == expected for rel in rels
        )

    @pytest.mark.parametrize(
        "instruction",
        [
            r'HYPERLINK "C:\docs\a.docx"',
            r'HYPERLINK "https://example.test/a""b"',
            r'HYPERLINK "https://example.test/a\q"',
            'HYPERLINK "https://example.test/trailing\\',
            'HYPERLINK "https://example.test/unclosed',
        ],
    )
    def it_skips_unsupported_or_malformed_quoted_escapes(self, instruction: str):
        content = package(document_parts(complex_field(instruction)))
        result = normalize(content, hyperlink_fields=True)
        assert result.report.succeeded
        assert result.report.skipped
        assert not result.report.changed
        assert result.content is content

    def it_decodes_escaped_quotes_in_field_switch_values(self):
        body = complex_field(r'HYPERLINK "https://example.test" \o "say \"hello\""')
        result = normalize(package(document_parts(body)), hyperlink_fields=True)
        assert result.report.succeeded
        link = next(_root(result.content).iter(f"{{{W}}}hyperlink"))
        assert link.get(f"{{{W}}}tooltip") == 'say "hello"'

    @pytest.mark.parametrize(
        "body",
        [
            complex_field().replace(
                "<w:r><w:rPr><w:i/></w:rPr><w:t>world</w:t></w:r>",
                "<w:ins><w:r><w:t>world</w:t></w:r></w:ins>",
            ),
            complex_field().replace('<w:r><w:fldChar w:fldCharType="end"/></w:r>', ""),
            complex_field().replace(
                "<w:r><w:rPr><w:i/></w:rPr><w:t>world</w:t></w:r>",
                '<w:bookmarkStart w:id="1" w:name="keep"/>'
                '<w:r><w:t>world</w:t></w:r><w:bookmarkEnd w:id="1"/>',
            ),
            complex_field('HYPERLINK "https://example.test" \\unknown value'),
            complex_field().replace(
                "<w:r><w:rPr><w:i/></w:rPr><w:t>world</w:t></w:r>",
                complex_field().removeprefix("<w:p>").removesuffix("</w:p>"),
            ),
            complex_field().replace(
                'w:fldCharType="begin"', 'w:fldCharType="begin" w:fldLock="true"'
            ),
            complex_field().replace(
                '<w:fldChar w:fldCharType="begin"/>',
                '<w:fldChar w:fldCharType="begin"/><w:t>must not disappear</w:t>',
            ),
        ],
    )
    def it_skips_ambiguous_fields_without_modifying_bytes(self, body: str):
        content = package(document_parts(body))
        result = normalize(content, hyperlink_fields=True)
        assert result.report.succeeded
        assert not result.report.changed
        assert result.report.skipped
        assert result.content == content

    def it_leaves_other_fields_untouched(self):
        content = package(document_parts(complex_field('TOC \\o "1-3"')))
        result = normalize(content, hyperlink_fields=True)
        assert result.report.succeeded
        assert not result.report.changed
        assert not result.report.skipped
        assert result.content is content

    def it_preserves_locked_simple_fields(self):
        body = (
            '<w:p><w:fldSimple w:fldLock="true" w:instr=\'HYPERLINK "https://example.test"\'>'
            "<w:r><w:t>label</w:t></w:r></w:fldSimple></w:p>"
        )
        content = package(document_parts(body))
        result = normalize(content, hyperlink_fields=True)
        assert result.report.succeeded
        assert result.report.skipped
        assert result.content is content

    def it_converts_multiple_fields_and_avoids_relationship_id_collisions(self):
        body = complex_field() + complex_field('HYPERLINK "https://example.test/other"')
        result = normalize(package(document_parts(body)), hyperlink_fields=True)
        assert result.report.succeeded
        ids = [
            link.get(f"{{{REL_ATTR}}}id")
            for link in _root(result.content).iter(f"{{{W}}}hyperlink")
        ]
        assert len(ids) == len(set(ids)) == 2
        assert all(change.target for change in result.report.changes)

    def it_applies_selection_without_embedding_application_rules(self):
        content = package(document_parts(complex_field()))
        seen: list[HyperlinkTarget] = []

        def select(target: HyperlinkTarget) -> bool:
            seen.append(target)
            return target.address == "https://different.test"

        result = normalize(content, hyperlink_fields=True, select_hyperlink=select)
        assert result.report.succeeded
        assert not result.report.changed
        assert result.report.skipped
        assert result.content is content
        assert seen == [HyperlinkTarget("word/document.xml", "https://example.test/doc")]


class DescribeHyperlinkRunMerging:
    @pytest.mark.parametrize("attribute", ["rsidR", "rsidRPr"])
    def it_merges_runs_with_different_session_ids_and_retains_the_first_id(self, attribute: str):
        body = (
            link_runs()
            .replace(
                '<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">',
                f'<w:r w:{attribute}="11111111"><w:rPr><w:b/></w:rPr><w:t xml:space="preserve">',
            )
            .replace(
                "<w:r><w:rPr><w:b/></w:rPr><w:t>world",
                f'<w:r w:{attribute}="22222222"><w:rPr><w:b/></w:rPr><w:t>world',
            )
        )
        result = normalize(package(document_parts(body)), hyperlink_runs=True)
        assert result.report.succeeded
        assert result.report.changed
        link = next(_root(result.content).iter(f"{{{W}}}hyperlink"))
        runs = link.findall(f"{{{W}}}r")
        assert len(runs) == 1
        assert runs[0].get(f"{{{W}}}{attribute}") == "11111111"
        assert runs[0].findtext(f"{{{W}}}t") == "Hello world"
        again = normalize(result.content, hyperlink_runs=True)
        assert not again.report.changed
        assert again.content == result.content

    def it_keeps_deletion_session_id_differences_as_a_merge_boundary(self):
        body = link_runs().replace(
            "<w:r><w:rPr><w:b/></w:rPr><w:t>world",
            '<w:r w:rsidDel="22222222"><w:rPr><w:b/></w:rPr><w:t>world',
        )
        content = package(document_parts(body))
        result = normalize(content, hyperlink_runs=True)
        assert result.report.succeeded
        assert not result.report.changed
        assert result.content is content

    def it_merges_compatible_runs_without_losing_spaces_or_formatting(self):
        content = package(document_parts(link_runs()))
        result = normalize(content, hyperlink_runs=True)
        assert result.report.succeeded
        assert result.report.changed
        link = next(_root(result.content).iter(f"{{{W}}}hyperlink"))
        runs = link.findall(f"{{{W}}}r")
        assert len(runs) == 1
        texts = runs[0].findall(f"{{{W}}}t")
        assert len(texts) == 1
        assert texts[0].text == "Hello world"
        assert texts[0].get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
        assert runs[0].find(f"{{{W}}}rPr/{{{W}}}b") is not None
        again = normalize(result.content, hyperlink_runs=True)
        assert not again.report.changed
        assert again.content == result.content

    @pytest.mark.parametrize(
        "body",
        [
            link_runs("<w:i/>"),
            link_runs().replace("<w:t>world</w:t>", "<w:tab/><w:t>world</w:t>"),
            link_runs().replace(
                "<w:r><w:rPr><w:b/></w:rPr><w:t>world",
                '<w:r w:customAttribute="different"><w:rPr><w:b/></w:rPr><w:t>world',
            ),
            link_runs().replace(
                "</w:r><w:r>", '</w:r><w:bookmarkStart w:id="1" w:name="keep"/><w:r>'
            ),
            "<w:p><w:ins>"
            + link_runs().removeprefix("<w:p>").removesuffix("</w:p>")
            + "</w:ins></w:p>",
        ],
    )
    def it_preserves_different_formats_nontext_content_and_markup_boundaries(self, body: str):
        content = package(document_parts(body))
        result = normalize(content, hyperlink_runs=True)
        assert result.report.succeeded
        assert not result.report.changed
        assert result.report.skipped
        assert result.content == content

    def it_applies_selection_to_existing_hyperlinks(self):
        content = package(document_parts(link_runs()))
        result = normalize(content, hyperlink_runs=True, select_hyperlink=lambda target: False)
        assert result.report.succeeded
        assert result.content is content


class DescribeNormalizationPackageContracts:
    def it_is_disabled_by_default(self):
        content = package(document_parts(complex_field() + link_runs()))
        result = normalize(content)
        assert result.content is content
        assert not result.report.changed

    def it_composes_conversion_and_merging(self):
        body = complex_field().replace("<w:i/>", "<w:b/>")
        result = normalize(
            package(document_parts(body)), hyperlink_fields=True, hyperlink_runs=True
        )
        assert result.report.succeeded
        link = next(_root(result.content).iter(f"{{{W}}}hyperlink"))
        assert len(link.findall(f"{{{W}}}r")) == 1
        assert {change.code for change in result.report.changes} >= {
            "converted_hyperlink_field",
            "merged_hyperlink_runs",
        }

    def it_preserves_unrelated_members_exactly(self):
        parts = document_parts(complex_field())
        parts["word/media/opaque.bin"] = b"\x00binary data\xff" * 1000
        parts["customXml/private.xml"] = b'<payload xmlns="urn:example:private">keep</payload>'
        content = package(parts)
        result = normalize(content, hyperlink_fields=True)
        assert result.report.succeeded
        after = unzip(result.content)
        changed = {name for name in parts if parts[name] != after[name]}
        assert changed == {"word/document.xml", "word/_rels/document.xml.rels"}

    def it_rolls_back_when_input_or_output_is_invalid(self):
        parts = document_parts(complex_field())
        parts["customXml/broken.xml"] = b"<broken>"
        content = package(parts)
        result = normalize(content, hyperlink_fields=True)
        assert not result.report.succeeded
        assert result.content is content
        assert not result.report.changed

    def it_restores_streams_and_keeps_files_unchanged(self, tmp_path: Path):
        content = package(document_parts(complex_field()))
        stream = io.BytesIO(content)
        stream.seek(7)
        assert normalize(stream, hyperlink_fields=True).report.succeeded
        assert stream.tell() == 7
        assert not stream.closed
        assert stream.getvalue() == content
        path = tmp_path / "source.docx"
        path.write_bytes(content)
        assert normalize(path, hyperlink_fields=True).report.succeeded
        assert path.read_bytes() == content

    def it_normalizes_related_header_story_parts(self):
        parts = document_parts()
        rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
        etree.SubElement(
            rels, f"{{{R}}}Relationship", Id="header99", Type=RT + "header", Target="header99.xml"
        )
        parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
        parts["word/header99.xml"] = f'<w:hdr xmlns:w="{W}">{complex_field()}</w:hdr>'.encode()
        types = etree.fromstring(parts["[Content_Types].xml"])
        etree.SubElement(
            types,
            f"{{{CT}}}Override",
            PartName="/word/header99.xml",
            ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml",
        )
        parts["[Content_Types].xml"] = etree.tostring(types)
        result = normalize(package(parts), hyperlink_fields=True)
        assert result.report.succeeded
        assert any(change.partname == "word/header99.xml" for change in result.report.changes)
        assert "word/_rels/header99.xml.rels" in unzip(result.content)

    def it_adds_content_types_for_new_relationship_members_without_a_rels_default(self):
        parts = document_parts(complex_field())
        types = etree.fromstring(parts["[Content_Types].xml"])
        for node in list(types):
            if node.get("Extension") == "rels":
                types.remove(node)
        del parts["word/_rels/document.xml.rels"]
        etree.SubElement(
            types,
            f"{{{CT}}}Override",
            PartName="/_rels/.rels",
            ContentType="application/vnd.openxmlformats-package.relationships+xml",
        )
        parts["[Content_Types].xml"] = etree.tostring(types)
        result = normalize(package(parts), hyperlink_fields=True)
        assert result.report.succeeded
        added = next(
            part
            for part in result.report.after.inventory.parts
            if part.partname == "word/_rels/document.xml.rels"
        )
        assert added.content_type == "application/vnd.openxmlformats-package.relationships+xml"
        assert any(
            change.code == "added_relationship_content_type" for change in result.report.changes
        )

    def it_propagates_selector_errors_without_changing_inputs(self):
        content = package(document_parts(complex_field()))
        stream = io.BytesIO(content)
        stream.seek(11)

        def failing_selector(target: HyperlinkTarget) -> bool:
            raise RuntimeError("Selector programming error")

        with pytest.raises(RuntimeError, match="Selector programming error"):
            normalize(stream, hyperlink_fields=True, select_hyperlink=failing_selector)
        assert stream.tell() == 11
        assert not stream.closed
        assert stream.getvalue() == content
