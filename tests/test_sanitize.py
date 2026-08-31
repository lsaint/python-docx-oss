"""Integration tests for package-level document sanitization."""

from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZipFile

from lxml import etree

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml

COMMENTS_FIXTURE = Path(__file__).parent / "fixtures/comments/word-comments-rich.docx"


def it_is_noop_without_selected_operations():
    document = Document()

    report = document.sanitize()

    assert report.changed is False
    assert report.removed_parts == ()
    assert report.removed_relationships == ()


def it_sanitizes_personal_metadata_canonically_and_idempotently():
    document = Document()
    root = document.part.package._core_properties_part.element
    etree.SubElement(root, qn("dc:creator")).text = "duplicate author"
    etree.SubElement(root, qn("cp:lastModifiedBy")).text = "duplicate editor"
    etree.SubElement(root, qn("cp:revision")).text = "9"

    first_report = document.sanitize(personal_metadata=True)
    second_report = document.sanitize(personal_metadata=True)
    reopened = Document(document.to_bytes())

    assert first_report.scrubbed_properties == (
        "author",
        "last_modified_by",
        "revision",
    )
    assert first_report.changed is True
    assert second_report.changed is False
    assert reopened.core_properties.author == ""
    assert reopened.core_properties.last_modified_by == ""
    assert reopened.core_properties.revision == 1


def it_preserves_existing_comments_on_a_roundtrip_without_sanitization():
    document = Document(COMMENTS_FIXTURE)
    original_comment_count = len(document.comments)
    original_marker_count = _comment_marker_count(COMMENTS_FIXTURE.read_bytes())

    reopened = Document(document.to_bytes())
    package = reopened.to_bytes()

    assert original_comment_count > 0
    assert original_marker_count > 0
    assert len(reopened.comments) == original_comment_count
    assert _comment_marker_count(package) == original_marker_count


def it_removes_comment_markers_and_the_comment_part_tree():
    document = Document()
    run = document.add_paragraph().add_run("annotated")
    document.add_comment(run, text="review", author="Reviewer")
    header_paragraph = document.sections[0].header.add_paragraph()
    header_paragraph._p.append(OxmlElement("w:commentReference"))

    comments_part = document.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
    )
    comments_extended = Part(
        PackURI("/word/commentsExtended.xml"),
        "application/vnd.ms-word.commentsExtended+xml",
        b'<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"/>',
        document.part.package,
    )
    comments_part.relate_to(
        comments_extended,
        "http://schemas.microsoft.com/office/2011/relationships/commentsExtended",
    )

    report = document.sanitize(comments=True)
    package = document.to_bytes()
    second_report = document.sanitize(comments=True)

    assert report.removed_comment_markers == 4
    assert report.removed_parts == (
        "/word/comments.xml",
        "/word/commentsExtended.xml",
    )
    assert len(report.removed_relationships) == 2
    assert second_report.changed is False
    with ZipFile(io.BytesIO(package)) as archive:
        assert "word/comments.xml" not in archive.namelist()
        assert "word/commentsExtended.xml" not in archive.namelist()
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                xml = archive.read(name)
                assert b"commentRangeStart" not in xml
                assert b"commentRangeEnd" not in xml
                assert b"commentReference" not in xml
        assert b"CommentReference" not in archive.read("word/document.xml")


def it_removes_the_entire_web_extension_relationship_tree():
    document = Document()
    package = document.part.package
    taskpanes = Part(
        PackURI("/word/webextensions/taskpanes.xml"),
        "application/vnd.ms-office.webextensiontaskpanes+xml",
        b"<wetp:taskpanes xmlns:wetp='urn:test'/>",
        package,
    )
    extension = Part(
        PackURI("/word/webextensions/webextension1.xml"),
        "application/vnd.ms-office.webextension+xml",
        b"<we:webextension xmlns:we='urn:test'/>",
        package,
    )
    package.relate_to(
        taskpanes,
        "http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes",
    )
    taskpanes.relate_to(
        extension,
        "http://schemas.microsoft.com/office/2011/relationships/webextension",
    )

    report = document.sanitize(web_extensions=True)
    package_bytes = document.to_bytes()

    assert report.removed_parts == (
        "/word/webextensions/taskpanes.xml",
        "/word/webextensions/webextension1.xml",
    )
    assert len(report.removed_relationships) == 2
    with ZipFile(io.BytesIO(package_bytes)) as archive:
        assert not any(name.startswith("word/webextensions/") for name in archive.namelist())


def it_round_trips_cleanup_without_removing_application_content():
    document = Document(COMMENTS_FIXTURE)
    custom_xml = document.part.add_custom_xml_part(
        '<application xmlns="urn:python-docx-oss:application" version="1"/>'
    )
    business_relationship_id = document.part.relate_to(
        "https://example.invalid/business/object-id",
        RT.HYPERLINK,
        is_external=True,
    )
    package = document.part.package
    taskpanes = Part(
        PackURI("/word/webextensions/taskpanes.xml"),
        "application/vnd.ms-office.webextensiontaskpanes+xml",
        b"<wetp:taskpanes xmlns:wetp='urn:test'/>",
        package,
    )
    extension = Part(
        PackURI("/word/webextensions/webextension1.xml"),
        "application/vnd.ms-office.webextension+xml",
        b"<we:webextension xmlns:we='urn:test'/>",
        package,
    )
    package.relate_to(
        taskpanes,
        "http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes",
    )
    taskpanes.relate_to(
        extension,
        "http://schemas.microsoft.com/office/2011/relationships/webextension",
    )

    report = document.sanitize(comments=True, web_extensions=True)
    output = document.to_bytes()
    reopened = Document(output)

    assert report.removed_comment_markers > 0
    assert str(custom_xml.partname) not in report.removed_parts
    assert reopened.part.custom_xml_parts[0].tag == ("{urn:python-docx-oss:application}application")
    assert reopened.part.rels[business_relationship_id].target_ref == (
        "https://example.invalid/business/object-id"
    )
    assert reopened.sanitize(comments=True, web_extensions=True).changed is False
    with ZipFile(io.BytesIO(output)) as archive:
        names = archive.namelist()
        assert "customXml/item1.xml" in names
        assert "word/comments.xml" not in names
        assert not any(name.startswith("word/webextensions/") for name in names)


def it_removes_formatting_revisions_and_keeps_current_formatting():
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("text")
    run.bold = True
    table = document.add_table(1, 1)

    run._r.get_or_add_rPr().append(OxmlElement("w:rPrChange"))
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:pPrChange"))
    table._tbl.tblPr.append(OxmlElement("w:tblPrChange"))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:trPrChange"))
    table.cell(0, 0)._tc.get_or_add_tcPr().append(OxmlElement("w:tcPrChange"))
    document.sections[0]._sectPr.append(OxmlElement("w:sectPrChange"))

    report = document.sanitize(formatting_revisions=True)
    reopened = Document(document.to_bytes())

    assert report.removed_formatting_revisions == 6
    assert reopened.paragraphs[0].runs[0].bold is True
    assert reopened.sanitize(formatting_revisions=True).changed is False


def it_counts_only_outermost_nested_formatting_revisions():
    document = Document()
    paragraph = document.add_paragraph("text")
    outer_change = OxmlElement("w:pPrChange")
    nested_properties = OxmlElement("w:pPr")
    nested_properties.append(OxmlElement("w:rPrChange"))
    outer_change.append(nested_properties)
    paragraph._p.get_or_add_pPr().append(outer_change)

    report = document.sanitize(formatting_revisions=True)

    assert report.removed_formatting_revisions == 1


def _comment_marker_count(package: bytes) -> int:
    marker_tags = (
        qn("w:commentRangeStart"),
        qn("w:commentRangeEnd"),
        qn("w:commentReference"),
    )
    count = 0
    with ZipFile(io.BytesIO(package)) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            root = parse_xml(archive.read(name))
            count += sum(len(root.findall(".//" + tag)) for tag in marker_tags)
    return count
